from __future__ import annotations
import os, sqlite3, hashlib, hmac, secrets, json, base64, re, io, time, zipfile, smtplib, ssl
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Optional
from urllib.parse import urlparse
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from fastapi import FastAPI, HTTPException, Depends, UploadFile, File, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def load_env_file(path: Path) -> None:
    """Load simple KEY=VALUE entries from .env without overriding real environment variables."""
    if not path.is_file():
        return
    try:
        for raw_line in path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip().strip("\\\"").strip("'")
            if key and key not in os.environ:
                os.environ[key] = value
    except OSError:
        # Environment variables can still be supplied by the shell/service.
        pass


load_env_file(ROOT / ".env")

_db_value = Path(os.getenv("DATABASE_PATH", str(ROOT / "data" / "ai_review.db")))
DB_PATH = _db_value if _db_value.is_absolute() else (ROOT / _db_value).resolve()
DB_PATH.parent.mkdir(parents=True, exist_ok=True)
SECRET = os.getenv("APP_SECRET", "local-demo-secret-change-before-deploy").encode()
DEMO_MODE = os.getenv("DEMO_MODE", "true").lower() == "true"
app = FastAPI(title='AI Review Automation System', version='1.0.0', description='Student placement and internship review workflow API')
app.add_middleware(CORSMiddleware, allow_origins=['http://127.0.0.1:8000','http://localhost:8000'], allow_credentials=True, allow_methods=['*'], allow_headers=['*'])


def connect():
    con = sqlite3.connect(DB_PATH, timeout=15)
    con.row_factory = sqlite3.Row
    con.execute('PRAGMA foreign_keys=ON')
    return con

def now(): return datetime.now(timezone.utc).isoformat()
def hash_password(password: str, salt: Optional[bytes]=None) -> str:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac('sha256', password.encode(), salt, 240_000)
    return base64.b64encode(salt + digest).decode()
def verify_password(password: str, stored: str) -> bool:
    try:
        raw=base64.b64decode(stored); return hmac.compare_digest(raw[16:], hashlib.pbkdf2_hmac('sha256', password.encode(), raw[:16], 240_000))
    except Exception: return False
def token_for(user):
    payload = json.dumps({'id':user['id'],'role':user['role'],'exp':int(time.time())+12*3600}, separators=(',',':')).encode()
    body=base64.urlsafe_b64encode(payload).decode().rstrip('=')
    sig=base64.urlsafe_b64encode(hmac.new(SECRET,body.encode(),hashlib.sha256).digest()).decode().rstrip('=')
    return body+'.'+sig
def decode_token(token):
    try:
        body,sig=token.split('.',1); expected=base64.urlsafe_b64encode(hmac.new(SECRET,body.encode(),hashlib.sha256).digest()).decode().rstrip('=')
        if not hmac.compare_digest(sig,expected): return None
        payload=json.loads(base64.urlsafe_b64decode(body+'='*(-len(body)%4)))
        return payload if payload['exp'] > time.time() else None
    except Exception: return None

def init_db():
    with connect() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, email TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('admin','reviewer','student')), student_id TEXT UNIQUE, display_name TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS students(id INTEGER PRIMARY KEY, student_id TEXT UNIQUE NOT NULL, name TEXT NOT NULL, email TEXT NOT NULL, phone TEXT, internship TEXT, task TEXT, number_of_reviews INTEGER DEFAULT 0, repository_url TEXT, submission_status TEXT DEFAULT 'Not Submitted', source TEXT DEFAULT 'import', created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS tasks(id INTEGER PRIMARY KEY, internship TEXT NOT NULL, task_name TEXT NOT NULL, problem_statement TEXT, abstract TEXT, working_flow TEXT, technologies TEXT, tools TEXT, deliverables TEXT, UNIQUE(internship,task_name));
        CREATE TABLE IF NOT EXISTS schedules(id INTEGER PRIMARY KEY, student_id TEXT NOT NULL, student_name TEXT, email TEXT, review_number INTEGER, review_date TEXT, review_time TEXT, meeting_link TEXT, status TEXT DEFAULT 'Scheduled', feedback_sent TEXT, timezone TEXT DEFAULT 'Asia/Kolkata', FOREIGN KEY(student_id) REFERENCES students(student_id));
        CREATE TABLE IF NOT EXISTS reviews(id INTEGER PRIMARY KEY, student_id TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'Awaiting submission', repository_url TEXT, draft_json TEXT, published_json TEXT, reviewer_id INTEGER, rejection_reason TEXT, updated_at TEXT NOT NULL, FOREIGN KEY(student_id) REFERENCES students(student_id));
        CREATE TABLE IF NOT EXISTS import_history(id INTEGER PRIMARY KEY, filename TEXT, imported_at TEXT NOT NULL, rows_seen INTEGER, created_count INTEGER, updated_count INTEGER, errors_json TEXT);
        CREATE TABLE IF NOT EXISTS audit_log(id INTEGER PRIMARY KEY, actor_id INTEGER, action TEXT NOT NULL, details TEXT, created_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_students_email ON students(email); CREATE INDEX IF NOT EXISTS idx_schedules_date ON schedules(review_date);
        ''')
        for email,password,role,name in [('admin@demo.local','AdminDemo123!','admin','Company Administrator'),('reviewer@demo.local','ReviewerDemo123!','reviewer','Company Reviewer')]:
            if not c.execute('SELECT id FROM users WHERE email=?',(email,)).fetchone(): c.execute('INSERT INTO users(email,password_hash,role,display_name,created_at) VALUES(?,?,?,?,?)',(email,hash_password(password),role,name,now()))
        # Use workbook's generated sample records for a useful first-run demo. Import is idempotent.
        if c.execute('SELECT COUNT(*) n FROM students').fetchone()['n']==0:
            sample=ROOT/'data'/'sample_workbook.xlsx'
            if sample.exists(): _import_workbook(c, sample.read_bytes(), sample.name, record_history=False)
        for row in c.execute('SELECT student_id,name,email FROM students').fetchall():
            if not c.execute('SELECT id FROM users WHERE student_id=?',(row['student_id'],)).fetchone():
                # Student demo sign-in is a common demo password; force replacement before production use.
                c.execute('INSERT OR IGNORE INTO users(email,password_hash,role,student_id,display_name,created_at) VALUES(?,?,?,?,?,?)',(row['email'],hash_password('StudentDemo123!'),'student',row['student_id'],row['name'],now()))

def _s(v):
    if v is None: return None
    if isinstance(v, datetime): return v.date().isoformat()
    return str(v).strip() or None

def _import_workbook(c, content: bytes, filename='upload.xlsx', record_history=True):
    try: wb=load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as e: raise HTTPException(400, 'Could not read workbook. Check that it is a valid .xlsx file.') from e
    errors=[]; seen=0; created=0; updated=0; students_sheet=None
    for s in wb.worksheets:
        if s.title.strip().lower()=='students': students_sheet=s
    if students_sheet is None:
        raise HTTPException(422, 'Required worksheet “Students” was not found. Available sheets: '+', '.join(wb.sheetnames))
    rows=students_sheet.iter_rows(values_only=True); headers=[str(x).strip() if x is not None else '' for x in next(rows,())]; mapping={h:i for i,h in enumerate(headers) if h}
    required=['Student_ID','Student_Name','Email']
    missing=[x for x in required if x not in mapping]
    if missing: raise HTTPException(422, 'Missing required columns in Students worksheet: '+', '.join(missing))
    student_seen=set(); email_seen=set()
    for ix,vals in enumerate(rows, start=2):
        if not any(v is not None and str(v).strip() for v in vals): continue
        seen+=1
        def get(k): return vals[mapping[k]] if k in mapping and mapping[k] < len(vals) else None
        sid=_s(get('Student_ID')); name=_s(get('Student_Name')); email=_s(get('Email'))
        why=[]
        if not sid: why.append('Student_ID is required')
        if not name: why.append('Student_Name is required')
        if not email or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',email): why.append('Email is missing or invalid')
        if sid in student_seen: why.append('Duplicate Student_ID within workbook')
        if email and email.lower() in email_seen: why.append('Duplicate email within workbook')
        if why:
            errors.append({'worksheet':'Students','row':ix,'student_id':sid,'field':'; '.join(why)}); continue
        student_seen.add(sid); email_seen.add(email.lower())
        values=(sid,name,email,_s(get('Phone_Number')),_s(get('Internship_Name')),_s(get('Task_Name')),int(get('Number_of_Reviews') or 0),_s(get('GitHub_Repository_URL')),_s(get('Submission_Status')) or 'Not Submitted')
        existing=c.execute('SELECT id FROM students WHERE student_id=?',(sid,)).fetchone()
        if existing:
            c.execute('UPDATE students SET name=?,email=?,phone=?,internship=?,task=?,number_of_reviews=?,repository_url=COALESCE(?,repository_url),submission_status=? WHERE student_id=?',(name,email,values[3],values[4],values[5],values[6],values[7],values[8],sid)); updated+=1
        else:
            c.execute('INSERT INTO students(student_id,name,email,phone,internship,task,number_of_reviews,repository_url,submission_status,source,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(*values,'import',now())); created+=1
        # Ensure the linked student account exists without overwriting an existing password.
        c.execute('INSERT OR IGNORE INTO users(email,password_hash,role,student_id,display_name,created_at) VALUES(?,?,?,?,?,?)',(email,hash_password('StudentDemo123!'),'student',sid,name,now()))
    # Task briefs: map by headers, preserving all available fields.
    for sheet in wb.worksheets:
        if sheet.title.strip().lower()=='task_brief':
            it=sheet.iter_rows(values_only=True); hs=[str(x).strip() if x is not None else '' for x in next(it,())]; mp={h:i for i,h in enumerate(hs) if h}
            for vals in it:
                def tg(k): return _s(vals[mp[k]]) if k in mp and mp[k]<len(vals) else None
                internship,task=tg('Internship_Name'),tg('Task_Name')
                if internship and task:
                    c.execute('INSERT INTO tasks(internship,task_name,problem_statement,abstract,working_flow,technologies,tools,deliverables) VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(internship,task_name) DO UPDATE SET problem_statement=excluded.problem_statement,abstract=excluded.abstract,working_flow=excluded.working_flow,technologies=excluded.technologies,tools=excluded.tools,deliverables=excluded.deliverables',(internship,task,tg('Problem_Statement'),tg('Abstract'),tg('Working_Flow'),tg('Suggested_Technologies'),tg('Suggested_Tools'),tg('Expected_Deliverables')))
    # Schedule import uses the dedicated schedule sheet if present; skip duplicate row inserts by key.
    for sheet in wb.worksheets:
        if sheet.title.strip().lower()=='review_schedule':
            it=sheet.iter_rows(values_only=True); hs=[str(x).strip() if x is not None else '' for x in next(it,())]; mp={h:i for i,h in enumerate(hs) if h}
            for vals in it:
                def sg(k): return _s(vals[mp[k]]) if k in mp and mp[k]<len(vals) else None
                sid=sg('Student_ID'); rn=sg('Review_Number')
                if not sid or not rn: continue
                date=sg('Review_Date'); rtime=sg('Review_Time'); link=sg('Google_Meet_Link')
                if not c.execute('SELECT id FROM schedules WHERE student_id=? AND review_number=? AND review_date=?',(sid,int(float(rn)),date)).fetchone():
                    c.execute('INSERT INTO schedules(student_id,student_name,email,review_number,review_date,review_time,meeting_link,status,feedback_sent) VALUES(?,?,?,?,?,?,?,?,?)',(sid,sg('Student_Name'),sg('Email'),int(float(rn)),date,rtime,link,sg('Status') or 'Scheduled',sg('Feedback_Sent')))
    if record_history:
        c.execute('INSERT INTO import_history(filename,imported_at,rows_seen,created_count,updated_count,errors_json) VALUES(?,?,?,?,?,?)',(filename,now(),seen,created,updated,json.dumps(errors)))
    return {'filename':filename,'worksheets':wb.sheetnames,'rows_seen':seen,'created':created,'updated':updated,'errors':errors,'error_count':len(errors),'tasks_imported':c.execute('SELECT COUNT(*) n FROM tasks').fetchone()['n'],'schedules_imported':c.execute('SELECT COUNT(*) n FROM schedules').fetchone()['n']}

init_db()

class LoginBody(BaseModel): email: str; password: str
class SubmissionBody(BaseModel): repository_url: str
class ReviewEdit(BaseModel): feedback: dict
class RejectBody(BaseModel): reason: str = Field(min_length=3,max_length=1000)

async def current_user(authorization: Optional[str]=Header(default=None)):
    if not authorization or not authorization.startswith('Bearer '): raise HTTPException(401,'Please sign in to continue.')
    data=decode_token(authorization[7:])
    if not data: raise HTTPException(401,'Session expired. Please sign in again.')
    with connect() as c: row=c.execute('SELECT id,email,role,student_id,display_name,active FROM users WHERE id=?',(data['id'],)).fetchone()
    if not row or not row['active']: raise HTTPException(401,'Account is unavailable.')
    return dict(row)
def require(*roles):
    async def dep(user=Depends(current_user)):
        if user['role'] not in roles: raise HTTPException(403,'You do not have permission to perform this action.')
        return user
    return dep

def audit(c, actor, action, details=''):
    c.execute('INSERT INTO audit_log(actor_id,action,details,created_at) VALUES(?,?,?,?)',(actor.get('id') if actor else None,action,details[:1000],now()))

def _send_review_email(recipient: str, student_name: str, feedback: dict, next_review: dict | None = None) -> None:
    host = os.getenv("SMTP_HOST", "").strip()
    username = os.getenv("SMTP_USERNAME", "").strip()
    password = "".join(os.getenv("SMTP_PASSWORD", "").split())
    sender = os.getenv("EMAIL_FROM", "").strip() or os.getenv("SMTP_FROM_EMAIL", "").strip() or username
    if not all((host, username, password, sender)):
        raise ValueError("SMTP_HOST, SMTP_USERNAME, SMTP_PASSWORD, and EMAIL_FROM must be configured.")

    lines = [
        f"Hello {student_name},",
        "",
        "Your project review has been approved and published.",
        "",
        f"Summary: {feedback.get('summary', 'Not provided')}",
        "",
        f"Overall assessment: {feedback.get('overall_assessment', 'Not provided')}",
    ]
    if feedback.get("score") is not None:
        lines.extend(("", f"Score: {feedback['score']}/100"))
    for title, key in (
        ("Strengths", "strengths"),
        ("Areas to improve", "improvements"),
        ("Recommended next steps", "recommendations"),
    ):
        items = feedback.get(key) or []
        if items:
            lines.extend(("", f"{title}:"))
            lines.extend(f"- {item}" for item in items)
    if feedback.get("repository_url"):
        lines.extend(("", f"Repository: {feedback['repository_url']}"))
    lines.extend(("", "Next review:"))
    if next_review:
        lines.append(f"Date: {next_review.get('review_date') or 'Not specified'}")
        lines.append(f"Time: {next_review.get('review_time') or 'Not specified'} ({next_review.get('timezone') or 'Asia/Kolkata'})")
        if next_review.get("meeting_link"):
            lines.append(f"Google Meet: {next_review['meeting_link']}")
        else:
            lines.append("Google Meet: No meeting link is available yet.")
    else:
        lines.append("No upcoming review is currently scheduled.")
    lines.extend(("", "Sign in to ReviewFlow to view your complete feedback."))

    message = EmailMessage()
    message["Subject"] = "Your project review has been published"
    message["From"] = sender
    message["To"] = recipient
    message.set_content("\n".join(lines))

    port = int(os.getenv("SMTP_PORT", "587"))
    if port == 465:
        with smtplib.SMTP_SSL(host, port, timeout=20, context=ssl.create_default_context()) as server:
            server.login(username, password)
            server.send_message(message)
    else:
        with smtplib.SMTP(host, port, timeout=20) as server:
            server.ehlo()
            server.starttls(context=ssl.create_default_context())
            server.ehlo()
            server.login(username, password)
            server.send_message(message)

def _ai_provider_config():
    provider = os.getenv("AI_PROVIDER", "openai").strip().lower()
    if provider == "claude":
        provider = "anthropic"
    if provider not in ("openai", "anthropic", "gemini"):
        return provider, "", ""

    key_names = {
        "openai": ("OPENAI_API_KEY",),
        "anthropic": ("ANTHROPIC_API_KEY", "CLAUDE_API_KEY"),
        "gemini": ("GEMINI_API_KEY",),
    }
    model_names = {
        "openai": ("OPENAI_MODEL",),
        "anthropic": ("ANTHROPIC_MODEL", "CLAUDE_MODEL"),
        "gemini": ("GEMINI_MODEL",),
    }
    defaults = {"openai": "gpt-4.1-mini", "anthropic": "claude-sonnet-4-5", "gemini": "gemini-2.5-flash"}
    api_key = next((os.getenv(name, "").strip() for name in key_names[provider] if os.getenv(name, "").strip()), "")
    model = next((os.getenv(name, "").strip() for name in model_names[provider] if os.getenv(name, "").strip()), defaults[provider])
    return provider, api_key, model

@app.get('/api/v1/health')
def health():
    try:
        with connect() as c: c.execute('SELECT 1')
        db='ok'
    except Exception: db='unavailable'
    provider, api_key, _ = _ai_provider_config()
    ai_status = 'configured' if api_key else ('unsupported provider' if provider not in ('openai', 'anthropic', 'gemini') else 'not configured')
    email_configured = all(os.getenv(key, "").strip() for key in ("SMTP_HOST", "SMTP_USERNAME", "SMTP_PASSWORD", "EMAIL_FROM"))
    email_status = 'demo mode' if DEMO_MODE else ('configured' if email_configured else 'not configured')
    return {'status':'ok' if db=='ok' else 'degraded','database':db,'demo_mode':DEMO_MODE,'integrations':{'ai':ai_status,'github':'configured' if os.getenv('GITHUB_TOKEN') else 'public repositories only / token not configured','email':email_status}}

@app.post('/api/v1/auth/login')
def login(body: LoginBody):
    with connect() as c: user=c.execute('SELECT * FROM users WHERE lower(email)=lower(?)',(body.email.strip(),)).fetchone()
    if not user or not verify_password(body.password,user['password_hash']) or not user['active']: raise HTTPException(401,'Email or password is incorrect.')
    u=dict(user); return {'access_token':token_for(u),'token_type':'bearer','user':{'id':u['id'],'email':u['email'],'role':u['role'],'student_id':u['student_id'],'display_name':u['display_name']}}
@app.get('/api/v1/auth/me')
def me(user=Depends(current_user)): return user
@app.post('/api/v1/auth/logout')
def logout(user=Depends(current_user)): return {'message':'Signed out. Clear the local session token.'}

@app.get('/api/v1/dashboard/admin')
def admin_dashboard(user=Depends(require('admin','reviewer'))):
    with connect() as c:
        n=lambda q: c.execute(q).fetchone()[0]
        return {'metrics':{'students':n('SELECT COUNT(*) FROM students'),'active_students':n("SELECT COUNT(*) FROM students WHERE submission_status!='Completed'"),'assigned_projects':n("SELECT COUNT(*) FROM students WHERE task IS NOT NULL AND task!=''"),'submissions_received':n("SELECT COUNT(*) FROM students WHERE repository_url IS NOT NULL AND repository_url!=''"),'submissions_pending':n("SELECT COUNT(*) FROM students WHERE repository_url IS NULL OR repository_url=''"),'reviews_awaiting_approval':n("SELECT COUNT(*) FROM reviews WHERE status IN ('Draft ready','Awaiting reviewer approval')"),'reviews_approved':n("SELECT COUNT(*) FROM reviews WHERE status IN ('Approved','Published','Completed')"),'upcoming_meetings':n("SELECT COUNT(*) FROM schedules WHERE review_date>=date('now') AND status='Scheduled'")},'recent_students':[dict(r) for r in c.execute('SELECT student_id,name,internship,task,submission_status FROM students ORDER BY id DESC LIMIT 8')],'demo_mode':DEMO_MODE}
@app.get('/api/v1/dashboard/student')
def student_dashboard(user=Depends(require('student'))):
    with connect() as c:
        st=c.execute('SELECT * FROM students WHERE student_id=?',(user['student_id'],)).fetchone()
        if not st: raise HTTPException(404,'Student record was not found.')
        d=dict(st); schedules=[dict(r) for r in c.execute('SELECT * FROM schedules WHERE student_id=? ORDER BY review_date,review_number',(d['student_id'],))]
        rev=c.execute("SELECT * FROM reviews WHERE student_id=? AND status IN ('Published','Approved','Completed') ORDER BY id DESC LIMIT 1",(d['student_id'],)).fetchone()
        return {'student':d,'schedules':schedules,'feedback':json.loads(rev['published_json']) if rev and rev['published_json'] else None,'review_status':rev['status'] if rev else 'Not started'}

@app.get('/api/v1/students')
def students(q: str='', internship: str='', status: str='', limit: int=100, offset: int=0, user=Depends(require('admin','reviewer'))):
    limit=max(1,min(limit,250)); sql='SELECT * FROM students WHERE 1=1'; args=[]
    if q: sql+=' AND (name LIKE ? OR student_id LIKE ? OR email LIKE ?)'; args += [f'%{q}%']*3
    if internship: sql+=' AND internship=?'; args.append(internship)
    if status: sql+=' AND submission_status=?'; args.append(status)
    sql+=' ORDER BY student_id LIMIT ? OFFSET ?'; args += [limit,offset]
    with connect() as c: return {'items':[dict(r) for r in c.execute(sql,args)],'limit':limit,'offset':offset}
@app.get('/api/v1/students/{student_id}')
def student_detail(student_id: str, user=Depends(current_user)):
    if user['role']=='student' and user['student_id']!=student_id: raise HTTPException(403,'You can only access your own student record.')
    if user['role'] not in ('admin','reviewer','student'): raise HTTPException(403,'Forbidden')
    with connect() as c:
        row=c.execute('SELECT * FROM students WHERE student_id=?',(student_id,)).fetchone()
        if not row: raise HTTPException(404,'Student not found.')
        d=dict(row); d['schedules']=[dict(r) for r in c.execute('SELECT * FROM schedules WHERE student_id=? ORDER BY review_date',(student_id,))]
        d['task_brief']=dict(c.execute('SELECT * FROM tasks WHERE internship=? AND task_name=?',(d['internship'],d['task'])).fetchone() or {})
        return d
@app.get('/api/v1/tasks')
def tasks(user=Depends(current_user)):
    with connect() as c: return [dict(r) for r in c.execute('SELECT * FROM tasks ORDER BY internship,task_name')]
@app.get('/api/v1/schedules')
def schedules(user=Depends(current_user)):
    with connect() as c:
        if user['role']=='student': rows=c.execute('SELECT * FROM schedules WHERE student_id=? ORDER BY review_date',(user['student_id'],))
        else: rows=c.execute('SELECT * FROM schedules ORDER BY review_date,review_time LIMIT 1000')
        return [dict(r) for r in rows]

@app.post('/api/v1/imports/validate')
async def validate_import(file: UploadFile=File(...), user=Depends(require('admin'))):
    if not file.filename or Path(file.filename).suffix.lower() not in ('.xlsx','.xlsm'): raise HTTPException(400,'Upload an .xlsx or .xlsm workbook.')
    content=await file.read()
    if len(content)>10*1024*1024: raise HTTPException(413,'Workbook exceeds the 10 MB upload limit.')
    # Validate in a rollback-only transaction.
    with connect() as c:
        c.execute('SAVEPOINT preview')
        try: result=_import_workbook(c,content,file.filename,record_history=False)
        finally: c.execute('ROLLBACK TO preview'); c.execute('RELEASE preview')
    return result
@app.post('/api/v1/imports/confirm')
async def confirm_import(file: UploadFile=File(...), user=Depends(require('admin'))):
    if not file.filename or Path(file.filename).suffix.lower() not in ('.xlsx','.xlsm'): raise HTTPException(400,'Upload an .xlsx or .xlsm workbook.')
    content=await file.read()
    if len(content)>10*1024*1024: raise HTTPException(413,'Workbook exceeds the 10 MB upload limit.')
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        try:
            result=_import_workbook(c,content,Path(file.filename).name,record_history=True); audit(c,user,'workbook_import',json.dumps({'filename':file.filename,'created':result['created'],'updated':result['updated']})); c.commit()
        except Exception: c.rollback(); raise
    return result
@app.get('/api/v1/imports/history')
def import_history(user=Depends(require('admin'))):
    with connect() as c: return [dict(r) for r in c.execute('SELECT id,filename,imported_at,rows_seen,created_count,updated_count,errors_json FROM import_history ORDER BY id DESC LIMIT 50')]

@app.post('/api/v1/submissions')
def submit(body: SubmissionBody,user=Depends(require('student'))):
    parsed=urlparse(body.repository_url.strip())
    if parsed.scheme!='https' or parsed.hostname not in ('github.com','www.github.com') or not re.fullmatch(r'/[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+/?',parsed.path): raise HTTPException(422,'Enter a valid GitHub repository URL, for example https://github.com/owner/repository.')
    url='https://github.com/'+parsed.path.strip('/')
    with connect() as c:
        c.execute('UPDATE students SET repository_url=?,submission_status=? WHERE student_id=?',(url,'Submitted',user['student_id']))
        row=c.execute('SELECT id FROM reviews WHERE student_id=? ORDER BY id DESC LIMIT 1',(user['student_id'],)).fetchone()
        if row: c.execute("UPDATE reviews SET repository_url=?,status='Submitted',updated_at=? WHERE id=?",(url,now(),row['id']))
        else: c.execute("INSERT INTO reviews(student_id,status,repository_url,updated_at) VALUES(?,?,?,?)",(user['student_id'],'Submitted',url,now()))
        audit(c,user,'repository_submitted',user['student_id'])
    return {'status':'Submitted','repository_url':url,'analysis_status':'Not analyzed yet','message':'Repository URL saved. Repository contents have not yet been analyzed.'}

@app.get('/api/v1/reviews')
def reviews(user=Depends(current_user)):
    with connect() as c:
        if user['role']=='student': rows=c.execute('SELECT r.*,s.name student_name FROM reviews r JOIN students s ON s.student_id=r.student_id WHERE r.student_id=? AND r.status IN (\'Published\',\'Approved\',\'Completed\') ORDER BY r.id DESC',(user['student_id'],))
        else: rows=c.execute('SELECT r.*,s.name student_name FROM reviews r JOIN students s ON s.student_id=r.student_id ORDER BY r.updated_at DESC LIMIT 500')
        return [{**dict(r),'draft':json.loads(r['draft_json']) if r['draft_json'] else None,'published':json.loads(r['published_json']) if r['published_json'] else None} for r in rows]
def _github_repo_parts(repository_url: str):
    parsed = urlparse((repository_url or "").strip())
    if parsed.scheme != "https" or parsed.hostname not in ("github.com", "www.github.com"):
        raise HTTPException(422, "The saved repository URL is not a valid HTTPS GitHub URL.")
    parts = [part for part in parsed.path.strip("/").split("/") if part]
    if len(parts) != 2:
        raise HTTPException(422, "Use a repository URL like https://github.com/owner/repository.")
    owner, repo = parts
    repo = repo.removesuffix(".git")
    if not re.fullmatch(r"[A-Za-z0-9_.-]+", owner) or not re.fullmatch(r"[A-Za-z0-9_.-]+", repo):
        raise HTTPException(422, "The GitHub repository URL contains invalid characters.")
    return owner, repo


def _github_headers():
    headers = {
        "Accept": "application/vnd.github+json",
        "User-Agent": "ReviewFlow-AI-Review/1.0",
        "X-GitHub-Api-Version": "2022-11-28",
    }
    token = os.getenv("GITHUB_TOKEN", "").strip()
    if token:
        headers["Authorization"] = f"Bearer {token}"
    return headers


def _read_github_project(repository_url: str) -> dict:
    """Fetch a bounded snapshot of a public or token-accessible GitHub repository."""
    owner, repo = _github_repo_parts(repository_url)
    api_base = f"https://api.github.com/repos/{owner}/{repo}"
    try:
        request = Request(api_base, headers=_github_headers())
        with urlopen(request, timeout=20) as response:
            metadata = json.loads(response.read(1_000_000).decode("utf-8"))
        branch = metadata.get("default_branch") or "main"
        archive_url = f"{api_base}/zipball/{branch}"
        request = Request(archive_url, headers=_github_headers())
        with urlopen(request, timeout=30) as response:
            archive_bytes = response.read(12 * 1024 * 1024 + 1)
        if len(archive_bytes) > 12 * 1024 * 1024:
            raise HTTPException(413, "Repository archive is larger than the 12 MB review limit.")
        files = []
        total_chars = 0
        allowed_suffixes = {
            ".py", ".ipynb", ".md", ".txt", ".csv", ".json", ".yml", ".yaml",
            ".toml", ".ini", ".cfg", ".sql", ".js", ".ts", ".html", ".css",
            ".r", ".rmd", ".ipynb", ".sh"
        }
        ignored_parts = {
            ".git", "node_modules", ".venv", "venv", "__pycache__", "dist",
            "build", ".idea", ".vscode", "vendor"
        }
        with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
            for info in archive.infolist():
                if info.is_dir() or info.file_size > 180_000:
                    continue
                path_parts = Path(info.filename).parts
                if any(part in ignored_parts for part in path_parts):
                    continue
                relative = "/".join(path_parts[1:]) if len(path_parts) > 1 else info.filename
                suffix = Path(relative).suffix.lower()
                basename = Path(relative).name.lower()
                if suffix not in allowed_suffixes and basename not in ("dockerfile", "makefile", ".gitignore"):
                    continue
                if any(secret_word in basename for secret_word in (".env", "secret", "credential", "private_key")):
                    continue
                try:
                    raw = archive.read(info)
                    text_content = raw.decode("utf-8", errors="replace")
                except Exception:
                    continue
                text_content = text_content.replace("\x00", "")
                if not text_content.strip():
                    continue
                remaining = 120_000 - total_chars
                if remaining <= 0 or len(files) >= 35:
                    break
                text_content = text_content[:min(24_000, remaining)]
                files.append({"path": relative, "content": text_content})
                total_chars += len(text_content)
        if not files:
            raise HTTPException(422, "No supported text, source-code, notebook, or documentation files were found in this repository.")
        return {
            "owner": owner,
            "repository": repo,
            "default_branch": branch,
            "description": metadata.get("description") or "",
            "html_url": metadata.get("html_url") or f"https://github.com/{owner}/{repo}",
            "files": files,
        }
    except HTTPException:
        raise
    except HTTPError as exc:
        if exc.code == 404:
            raise HTTPException(422, "GitHub could not access this repository. Check the URL and make sure the repository is public, or configure a token with access to it.") from exc
        if exc.code == 401 or exc.code == 403:
            raise HTTPException(502, "GitHub denied repository access. Check repository visibility, token permissions, or GitHub API rate limits.") from exc
        raise HTTPException(502, f"GitHub returned an error while retrieving the repository (HTTP {exc.code}).") from exc
    except (URLError, TimeoutError, zipfile.BadZipFile, json.JSONDecodeError, OSError) as exc:
        raise HTTPException(502, "Could not retrieve the GitHub repository. Check the network connection and repository URL, then try again.") from exc


def _generate_ai_review(project: dict, task: str | None, internship: str | None) -> dict:
    provider, api_key, model = _ai_provider_config()
    if provider not in ("openai", "anthropic", "gemini"):
        raise HTTPException(503, "Unsupported AI_PROVIDER. Choose openai, claude, or gemini.")
    if not api_key:
        provider_name = "Claude" if provider == "anthropic" else provider.capitalize()
        key_name = {"openai": "OPENAI_API_KEY", "anthropic": "ANTHROPIC_API_KEY", "gemini": "GEMINI_API_KEY"}[provider]
        raise HTTPException(503, f"{provider_name} is not configured. Add {key_name} to the server environment and restart the backend.")
    if provider != "openai":
        return _generate_non_openai_review(provider, api_key, model, project, task, internship)

    project_files = []
    for item in project["files"]:
        project_files.append(
            f"FILE: {item['path']}\\nCONTENT:\\n{item['content']}"
        )
    project_context = "\\n\\n--- FILE SEPARATOR ---\\n\\n".join(project_files)
    if len(project_context) > 125_000:
        project_context = project_context[:125_000]

    system_prompt = (
        "You are a careful academic project reviewer. Evaluate only evidence visible in the supplied repository. "
        "Do not invent results, datasets, model metrics, or features. Be supportive, specific, and understandable "
        "to a student. If a detail cannot be verified from the files, say it is not clear from the submission. "
        "Return valid JSON only with these keys: summary (string), overall_assessment (string), strengths (array of strings), "
        "improvements (array of strings), recommendations (array of strings), files_reviewed (array of strings), "
        "limitations (array of strings), score (integer from 0 to 100 or null)."
    )
    user_prompt = (
        f"Internship: {internship or 'Not specified'}\\n"
        f"Assigned task: {task or 'Not specified'}\\n"
        f"Repository: {project['html_url']}\\n"
        f"Repository description: {project['description'] or 'Not provided'}\\n"
        f"Default branch: {project['default_branch']}\\n\\n"
        "Review the submitted project. Discuss relevance to the assigned task, code quality, data handling, methodology, "
        "evaluation, documentation, reproducibility, and practical next steps where applicable. Do not assume that an "
        "expected feature exists unless the files show it. Give concise, concrete feedback in plain English.\n\n"
        f"Repository files:\\n{project_context}"
    )
    payload = {
        "model": model,
        "temperature": 0.2,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    }
    request = Request(
        "https://api.openai.com/v1/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "User-Agent": "ReviewFlow-AI-Review/1.0",
        },
        method="POST",
    )
    try:
        with urlopen(request, timeout=90) as response:
            response_data = json.loads(response.read(2_000_000).decode("utf-8"))
        content = response_data["choices"][0]["message"]["content"]
        feedback = json.loads(content)
        required = ("summary", "overall_assessment", "strengths", "improvements", "recommendations")
        for key in required:
            if key not in feedback:
                raise ValueError(f"AI response is missing {key}")
        for key in ("strengths", "improvements", "recommendations", "files_reviewed", "limitations"):
            if key not in feedback or not isinstance(feedback[key], list):
                feedback[key] = []
        if not isinstance(feedback.get("summary"), str):
            feedback["summary"] = str(feedback["summary"])
        if not isinstance(feedback.get("overall_assessment"), str):
            feedback["overall_assessment"] = str(feedback["overall_assessment"])
        score = feedback.get("score")
        if not isinstance(score, int) or isinstance(score, bool) or not 0 <= score <= 100:
            feedback["score"] = None
        feedback["repository_url"] = project["html_url"]
        feedback["model"] = model
        feedback["provider"] = "openai"
        feedback["generated_at"] = now()
        return feedback
    except HTTPError as exc:
        # Do not expose provider response bodies or secrets to the browser.
        if exc.code == 401:
            raise HTTPException(502, "OpenAI rejected the API key. Check OPENAI_API_KEY and API account access.") from exc
        if exc.code == 429:
            raise HTTPException(502, "OpenAI rate limit or API quota was reached. Check your API billing and usage limits.") from exc
        raise HTTPException(502, f"OpenAI request failed (HTTP {exc.code}). Check the model name and API configuration.") from exc
    except (URLError, TimeoutError) as exc:
        raise HTTPException(502, "Could not connect to OpenAI. Check the server's internet connection and try again.") from exc
    except (KeyError, IndexError, TypeError, json.JSONDecodeError, ValueError) as exc:
        raise HTTPException(502, "OpenAI returned an unexpected response. Please try again.") from exc


def _generate_non_openai_review(provider: str, api_key: str, model: str, project: dict, task: str | None, internship: str | None) -> dict:
    project_files = [
        f"FILE: {item['path']}\\nCONTENT:\\n{item['content']}"
        for item in project["files"]
    ]
    project_context = "\\n\\n--- FILE SEPARATOR ---\\n\\n".join(project_files)[:125_000]
    system_prompt = (
        "You are a careful academic project reviewer. Evaluate only evidence visible in the supplied repository. "
        "Do not invent results, datasets, model metrics, or features. Be supportive, specific, and understandable "
        "to a student. If a detail cannot be verified from the files, say it is not clear from the submission. "
        "Return valid JSON only with these keys: summary (string), overall_assessment (string), strengths (array of strings), "
        "improvements (array of strings), recommendations (array of strings), files_reviewed (array of strings), "
        "limitations (array of strings), score (integer from 0 to 100 or null)."
    )
    user_prompt = (
        f"Internship: {internship or 'Not specified'}\\n"
        f"Assigned task: {task or 'Not specified'}\\n"
        f"Repository: {project['html_url']}\\n"
        f"Repository description: {project['description'] or 'Not provided'}\\n"
        f"Default branch: {project['default_branch']}\\n\\n"
        "Review the submitted project. Discuss relevance to the assigned task, code quality, data handling, methodology, "
        "evaluation, documentation, reproducibility, and practical next steps where applicable. Do not assume that an "
        "expected feature exists unless the files show it. Give concise, concrete feedback in plain English.\\n\\n"
        f"Repository files:\\n{project_context}"
    )
    if provider == "anthropic":
        url = "https://api.anthropic.com/v1/messages"
        payload = {
            "model": model,
            "max_tokens": 4096,
            "temperature": 0.2,
            "system": system_prompt + " Return valid JSON only.",
            "messages": [{"role": "user", "content": user_prompt}],
        }
        headers = {
            "x-api-key": api_key,
            "anthropic-version": "2023-06-01",
            "Content-Type": "application/json",
        }
    else:
        from urllib.parse import quote
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{quote(model, safe='')}:generateContent"
        payload = {
            "systemInstruction": {"parts": [{"text": system_prompt + " Return valid JSON only."}]},
            "contents": [{"role": "user", "parts": [{"text": user_prompt}]}],
            "generationConfig": {"temperature": 0.2, "responseMimeType": "application/json"},
        }
        headers = {"x-goog-api-key": api_key, "Content-Type": "application/json"}

    request = Request(url, data=json.dumps(payload).encode("utf-8"), headers=headers, method="POST")
    try:
        with urlopen(request, timeout=90) as response:
            response_data = json.loads(response.read(2_000_000).decode("utf-8"))
        if provider == "anthropic":
            content = "".join(block["text"] for block in response_data["content"] if block.get("type") == "text")
        else:
            content = "".join(part["text"] for part in response_data["candidates"][0]["content"]["parts"] if "text" in part)
        feedback = json.loads(content)
        required = ("summary", "overall_assessment", "strengths", "improvements", "recommendations")
        for key in required:
            if key not in feedback:
                raise ValueError(f"AI response is missing {key}")
        for key in ("strengths", "improvements", "recommendations", "files_reviewed", "limitations"):
            if key not in feedback or not isinstance(feedback[key], list):
                feedback[key] = []
        if not isinstance(feedback.get("summary"), str):
            feedback["summary"] = str(feedback["summary"])
        if not isinstance(feedback.get("overall_assessment"), str):
            feedback["overall_assessment"] = str(feedback["overall_assessment"])
        score = feedback.get("score")
        if not isinstance(score, int) or isinstance(score, bool) or not 0 <= score <= 100:
            feedback["score"] = None
        feedback["repository_url"] = project["html_url"]
        feedback["model"] = model
        feedback["provider"] = "claude" if provider == "anthropic" else provider
        feedback["generated_at"] = now()
        return feedback
    except HTTPError as exc:
        provider_name = "Claude" if provider == "anthropic" else "Gemini"
        if exc.code == 401 or exc.code == 403:
            raise HTTPException(502, f"{provider_name} rejected the API key or denied access. Check the key and API account access.") from exc
        if exc.code == 429:
            raise HTTPException(502, f"{provider_name} rate limit or API quota was reached. Check API billing and usage limits.") from exc
        raise HTTPException(502, f"{provider_name} request failed (HTTP {exc.code}). Check the model name and API configuration.") from exc
    except (URLError, TimeoutError) as exc:
        provider_name = "Claude" if provider == "anthropic" else "Gemini"
        raise HTTPException(502, f"Could not connect to {provider_name}. Check the server's internet connection and try again.") from exc
    except (KeyError, IndexError, TypeError, json.JSONDecodeError, ValueError) as exc:
        provider_name = "Claude" if provider == "anthropic" else "Gemini"
        raise HTTPException(502, f"{provider_name} returned an unexpected response. Please try again.") from exc


@app.post('/api/v1/reviews/{review_id}/analyze')
def analyze(review_id: int, user=Depends(require('admin', 'reviewer'))):
    with connect() as c:
        row = c.execute(
            'SELECT r.*, s.task, s.internship FROM reviews r '
            'JOIN students s ON s.student_id = r.student_id WHERE r.id = ?',
            (review_id,)
        ).fetchone()
        if not row:
            raise HTTPException(404, 'Review not found.')
        review = dict(row)
        if not review.get("repository_url"):
            raise HTTPException(422, "This review has no saved GitHub repository URL. Ask the student to submit it first.")

    project = _read_github_project(review["repository_url"])
    feedback = _generate_ai_review(project, review.get("task"), review.get("internship"))

    with connect() as c:
        current = c.execute('SELECT id FROM reviews WHERE id = ?', (review_id,)).fetchone()
        if not current:
            raise HTTPException(404, 'Review not found.')
        c.execute(
            "UPDATE reviews SET draft_json = ?, status = 'Draft ready', updated_at = ? WHERE id = ?",
            (json.dumps(feedback, ensure_ascii=False), now(), review_id)
        )
        audit(c, user, 'ai_review_generated', json.dumps({
            'review_id': review_id,
            'repository_url': project['html_url'],
            'model': feedback.get('model'),
        }))
    return {
        'status': 'Draft ready',
        'message': 'AI review generated successfully. A reviewer must approve it before the student can see it.',
        'review_id': review_id,
        'feedback': feedback,
    }
@app.patch('/api/v1/reviews/{review_id}/draft')
def edit_review(review_id: int, body: ReviewEdit,user=Depends(require('admin','reviewer'))):
    with connect() as c:
        r=c.execute('SELECT * FROM reviews WHERE id=?',(review_id,)).fetchone()
        if not r: raise HTTPException(404,'Review not found.')
        if r['status'] not in ('Draft ready','Awaiting reviewer approval','Changes requested'): raise HTTPException(409,'This review has no editable draft.')
        c.execute("UPDATE reviews SET draft_json=?,updated_at=? WHERE id=?",(json.dumps(body.feedback),now(),review_id)); audit(c,user,'review_draft_edited',str(review_id))
    return {'status':'Draft saved'}
@app.post('/api/v1/reviews/{review_id}/approve')
def approve_review(review_id:int,user=Depends(require('admin','reviewer'))):
    with connect() as c:
        r=c.execute(
            'SELECT r.*, s.student_id AS linked_student_id, s.email AS student_email, s.name AS student_name '
            'FROM reviews r JOIN students s ON s.student_id=r.student_id WHERE r.id=?',
            (review_id,)
        ).fetchone()
        if not r: raise HTTPException(404,'Review not found.')
        if not r['draft_json']: raise HTTPException(409,'There is no validated draft to approve.')
        next_review = c.execute(
            "SELECT review_date,review_time,meeting_link,timezone FROM schedules "
            "WHERE student_id=? AND date(review_date)>=date('now') "
            "AND (status IS NULL OR lower(status) IN ('scheduled','upcoming')) "
            "ORDER BY date(review_date),review_time,review_number LIMIT 1",
            (r['linked_student_id'],)
        ).fetchone()
        next_review = dict(next_review) if next_review else None
        c.execute("UPDATE reviews SET status='Approved',published_json=?,reviewer_id=?,updated_at=? WHERE id=?",(r['draft_json'],user['id'],now(),review_id)); audit(c,user,'review_approved',str(review_id))
    feedback = json.loads(r['draft_json'])
    email_status = 'not_sent'
    if DEMO_MODE:
        email_message = 'Email was not sent because DEMO_MODE is enabled.'
    else:
        try:
            _send_review_email(r['student_email'], r['student_name'], feedback, next_review)
            email_status = 'sent'
            email_message = f"Notification email sent to {r['student_email']}."
        except (OSError, smtplib.SMTPException, ValueError) as exc:
            email_status = 'failed'
            email_message = 'Feedback was published, but the notification email could not be sent. Check Gmail SMTP settings and credentials.'
    return {
        'status':'Approved',
        'email_status':email_status,
        'message':f"Feedback approved and published. {email_message}",
    }
@app.post('/api/v1/reviews/{review_id}/request-changes')
def reject_review(review_id:int,body:RejectBody,user=Depends(require('admin','reviewer'))):
    with connect() as c:
        r=c.execute('SELECT id FROM reviews WHERE id=?',(review_id,)).fetchone()
        if not r: raise HTTPException(404,'Review not found.')
        c.execute("UPDATE reviews SET status='Changes requested',rejection_reason=?,reviewer_id=?,updated_at=? WHERE id=?",(body.reason,user['id'],now(),review_id)); audit(c,user,'review_changes_requested',str(review_id))
    return {'status':'Changes requested'}

@app.get('/api/v1/reports/students.csv')
def export_students(user=Depends(require('admin'))):
    from fastapi.responses import StreamingResponse
    import csv
    out=io.StringIO(); writer=csv.writer(out); writer.writerow(['Student_ID','Name','Email','Internship','Task','Submission_Status'])
    with connect() as c:
        for r in c.execute('SELECT student_id,name,email,internship,task,submission_status FROM students ORDER BY student_id'):
            safe=[("'"+str(v) if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for v in r]; writer.writerow(safe)
    out.seek(0); return StreamingResponse(iter([out.getvalue()]),media_type='text/csv',headers={'Content-Disposition':'attachment; filename="student-progress.csv"'})

FRONTEND=ROOT/'frontend'
app.mount('/assets',StaticFiles(directory=FRONTEND),name='assets')
@app.get('/')
def home(): return FileResponse(FRONTEND/'index.html')
